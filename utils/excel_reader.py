import openpyxl
from datetime import datetime

class ExcelReader:

    def __init__(self, file_path):
        self.file_path = file_path

    def get_sheet_data(self, sheet_name):
        workbook = openpyxl.load_workbook(self.file_path)
        sheet = workbook[sheet_name]

        data = []
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = {}

            for key, value in zip(headers, row):

                # Convert datetime automatically
                if isinstance(value, datetime):
                    if key in ["PeriodBegin", "PeriodEnd"]:
                        value = value.strftime("%m-%Y")
                    elif key == "DueDate":
                        value = value.strftime("%d/%m/%Y")

                # Remove extra quotes from file path
                if key == "File_Path" and value:
                    value = str(value).replace('"', '').strip()

                row_data[key] = value

            data.append(row_data)

        return data
