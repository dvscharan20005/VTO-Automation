import openpyxl

def update_excel_status(file_path, sheet_name, row_index, status, column_name="Status"):
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook[sheet_name] 
    
    target_col = None
    
    # Search the first row to find the column index dynamically
    for col in range(1, sheet.max_column + 1):
        if sheet.cell(row=1, column=col).value == column_name:
            target_col = col
            break
            
    if target_col is None:
        raise ValueError(f"Column '{column_name}' not found in the Excel sheet.")

    sheet.cell(row=row_index, column=target_col).value = status
    
    workbook.save(file_path)